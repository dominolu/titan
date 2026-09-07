from __future__ import annotations

from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/dominolu/dev/titan")
OUT = ROOT / "docs/design/Titan_OpenCTP_Connector_技术方案.docx"


NAVY = "17365D"
BLUE = "2F5597"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F5F7F9"
MID_GRAY = "667085"
LIGHT_GRAY = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_east_asia(run, name="SimHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.28)
    sec.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)

    for name, size, before, after in (
        ("Heading 1", 17, 14, 8),
        ("Heading 2", 13, 11, 6),
        ("Heading 3", 11, 8, 4),
    ):
        st = styles[name]
        st.font.name = "Aptos"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Menlo"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    code.font.size = Pt(8.3)
    code.paragraph_format.left_indent = Inches(0.22)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.05
    p_pr = code._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GRAY)
    p_pr.append(shd)

    if "Caption Small" not in styles:
        cap = styles.add_style("Caption Small", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Caption Small"]
    cap.font.name = "Aptos"
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    cap.font.size = Pt(8.5)
    cap.font.color.rgb = RGBColor(102, 112, 133)
    cap.paragraph_format.space_after = Pt(5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_east_asia(run)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = paragraph.add_run(" 页")
    set_east_asia(run2)


def configure_header_footer(doc: Document):
    for section in doc.sections:
        hp = section.header.paragraphs[0]
        hp.text = "Titan OpenCTP Connector 技术方案"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in hp.runs:
            set_east_asia(run)
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(102, 112, 133)
        add_page_number(section.footer.paragraphs[0])
        for run in section.footer.paragraphs[0].runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(102, 112, 133)


def add_title(doc, text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_east_asia(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        set_east_asia(r)
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_para(doc, text="", bold_lead=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        set_east_asia(r)
        r2 = p.add_run(text[len(bold_lead):])
        set_east_asia(r2)
    else:
        r = p.add_run(text)
        set_east_asia(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_east_asia(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_east_asia(r)


def add_code(doc, code):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    for i, line in enumerate(code.strip("\n").splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Menlo"
        r.font.size = Pt(8.3)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        set_east_asia(run)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cell = cells[cidx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ridx % 2:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if cidx == 0 and len(str(value)) < 22:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_east_asia(run)
            run.font.size = Pt(font_size)
        if widths:
            for idx, width in enumerate(widths):
                cells[idx].width = Inches(width)
    if widths:
        for idx, width in enumerate(widths):
            header.cells[idx].width = Inches(width)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def page_break(doc):
    doc.add_page_break()


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("TITAN")
    r.font.name = "Aptos Display"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(47, 85, 151)

    add_title(doc, "Titan OpenCTP Connector 技术方案")
    p = doc.add_paragraph()
    r = p.add_run("自建薄 C++ Connector 与版本化 C ABI")
    set_east_asia(r)
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(102, 112, 133)
    p.paragraph_format.space_after = Pt(34)

    add_para(doc, "本方案定义 Titan 对接 CTP、CTP 股票期权、TTS、中泰 XTP、华鑫奇点、易盛 TAP 和量投 QDP 的统一工程路径。核心决策是保留 CTP API 与 SPI 在 C++ 边界内，通过小型、稳定、版本化的 C ABI 暴露给 Rust Sidecar，并以独立 Connector 进程隔离每个柜台的动态库、配置和故障。")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    for label, value in (
        ("文档状态", "技术设计基线"),
        ("适用范围", "Titan 实盘 Connector 扩展"),
        ("版本", "1.0"),
        ("日期", "2026 年 9 月 7 日"),
    ):
        line = p.add_run(f"{label}    {value}\n")
        set_east_asia(line)
        line.font.size = Pt(10)
    p.paragraph_format.line_spacing = 1.6


def build_document():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    page_break(doc)

    add_heading(doc, "文档控制", 1)
    add_table(doc, ["项目", "内容"], [
        ("目标读者", "Titan 架构、Connector、交易执行、运维和测试团队"),
        ("设计状态", "建议作为 PoC 与一期实施的架构基线，生产参数需在柜台联调后冻结"),
        ("核心依赖", "OpenCTP 兼容动态库、柜台厂商动态库、对应版本 CTP 头文件"),
        ("关键约束", "第三方动态库闭源、同名 SO 冲突、CTP C++ ABI、各柜台业务语义不完全一致"),
        ("仓库依据", "Titan 当前采用动态 Connector 插件、MarketPlugin、AccountPlugin 与 EventEngine"),
    ], widths=[1.35, 5.45])

    add_heading(doc, "变更记录", 2)
    add_table(doc, ["版本", "日期", "说明"], [
        ("1.0", "2026 09 07", "形成完整架构、接口、可靠性、测试、实施与验收方案"),
    ], widths=[0.8, 1.2, 4.8])

    add_heading(doc, "目录", 1)
    toc = [
        "1 执行摘要", "2 背景与设计判断", "3 目标与范围", "4 需求基线",
        "5 Titan 现状与集成约束", "6 总体架构", "7 组件设计", "8 进程与部署模型",
        "9 C ABI 设计", "10 IPC 与事件传输", "11 领域模型与标识", "12 生命周期与恢复",
        "13 柜台 Profile", "14 性能设计", "15 安全设计", "16 可观测性与运维",
        "17 测试与验收", "18 实施计划", "19 风险与决策", "20 附录",
    ]
    for item in toc:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        set_east_asia(r)
        r.font.size = Pt(9.5)

    page_break(doc)
    add_heading(doc, "1 执行摘要", 1)
    add_para(doc, "本方案建议 Titan 采用独立进程式 OpenCTP Connector。Titan 动态插件负责对接现有 MarketPlugin 和 AccountPlugin，并监管对应柜台 Sidecar；Rust Sidecar 负责 IPC、状态机和 C ABI 调用；C++ Bridge 使用原生 CTP C++ API 和 SPI 对接 OpenCTP SO。每个柜台实例拥有独立进程、动态库目录、流文件、日志和凭据。")
    add_para(doc, "这一结构保留了 ctp2rs 的主要效果，包括按版本生成或绑定 CTP 数据结构、加载不同柜台动态库、调用完整 CTP API 和接收 SPI 回调，同时避免在 Rust 中直接模拟 C++ 虚表。C++ 编译器负责类布局和虚函数调用，Rust 只信任经过版本校验的 C ABI 数据。")
    add_heading(doc, "1 1 核心决策", 2)
    add_table(doc, ["决策", "选择", "理由"], [
        ("语言边界", "C++ 到 C ABI 到 Rust", "把 C++ ABI 风险限制在 Bridge 内"),
        ("运行隔离", "一柜台实例一 Sidecar 进程", "隔离同名 SO、厂商依赖和崩溃"),
        ("集成方式", "Rust Connector 插件监管 Sidecar", "复用 Titan 当前动态插件、工厂和事件引擎"),
        ("版本策略", "头文件与 SO 成对固定", "防止函数签名、结构布局和虚表不匹配"),
        ("协议策略", "控制面可靠传输，数据面分级 QoS", "交易事件不得丢失，行情允许受控合并"),
        ("柜台差异", "Profile 显式建模", "避免把股票和期权特性隐藏在通用 CTP 字段中"),
    ], widths=[1.15, 1.7, 3.95])
    add_heading(doc, "1 2 建议交付顺序", 2)
    add_numbered(doc, [
        "完成 C ABI 与 Fake CTP 动态库，验证加载、回调、释放和故障注入。",
        "接入 TTS 标准 CTP 环境，打通行情、报单、撤单、查询和恢复。",
        "接入官方 CTP，完成生产认证、结算和交易日切换。",
        "按 TAP、QDP、TORA、TORA 股票期权、XTP 的风险顺序增加 Profile。",
        "每个柜台通过独立兼容、压力、故障恢复和实盘验收后再启用交易权限。",
    ])

    add_heading(doc, "2 背景与设计判断", 1)
    add_para(doc, "OpenCTP 通过提供与 CTPAPI 兼容的交易和行情动态库，将 TTS、CTP 股票期权、XTP、TORA、TAP、QDP 等柜台映射到 CTP 调用模型。该方式显著减少接入接口的数量，但兼容层仍受 CTP 公共模型、动态库 ABI 和柜台特例约束。")
    add_para(doc, "ctp2rs 已证明 Rust 可以通过 libloading、C++ 符号和虚表映射直接加载这些动态库。该项目覆盖多个 CTP 版本，并提供 OpenCTP 和股票期权示例。它也暴露出直接在 Rust 中承接 C++ ABI 的长期成本，包括硬编码 mangled symbol、unsafe Send 和 Sync、SPI 所有权、回调 panic 与版本升级风险。")
    add_heading(doc, "2 1 方案选择", 2)
    add_table(doc, ["方案", "优势", "主要问题", "结论"], [
        ("直接采用 ctp2rs", "开发快，API 覆盖广", "Rust 直接承担 C++ 虚表和生命周期风险", "用于参考和 PoC"),
        ("Rust 自建 C++ 虚表映射", "无额外 C++ Bridge", "重复 ctp2rs 的高风险工作", "不采用"),
        ("薄 C++ Bridge 加 C ABI", "ABI 边界清晰，易测试", "需要维护一层 C++", "采用"),
        ("全部柜台原生 API", "能力最完整", "开发和长期维护成本最高", "重点柜台后续可替换"),
    ], widths=[1.35, 1.7, 2.7, 1.05])
    add_heading(doc, "2 2 设计原则", 2)
    add_bullets(doc, [
        "第三方柜台库不得进入 Titan Core 进程。",
        "C++ 对象不得跨越 C ABI，Rust 只持有不透明句柄和 POD 数据。",
        "头文件、OpenCTP SO、厂商 SO、配置模板和校验摘要组成不可拆分的运行包。",
        "交易命令和交易回报使用可靠有序通道，行情使用独立的受控背压通道。",
        "通用模型无法表达的柜台字段必须进入版本化扩展区，不能静默丢弃。",
        "任何未知状态都通过查询与对账收敛，不以本地推测替代柜台事实。",
    ])

    add_heading(doc, "3 目标与范围", 1)
    add_heading(doc, "3 1 建设目标", 2)
    add_bullets(doc, [
        "为 Titan 提供统一的国内期货、股票和股票期权实盘 Connector 入口。",
        "在不修改 Titan Core 的前提下，通过现有 Connector 插件加载机制接入新柜台。",
        "将 CTP C++ ABI、OpenCTP 兼容库和厂商依赖隔离在可独立重启的进程中。",
        "建立可审计的订单标识、回报去重、断线恢复、持仓资金对账与交易日切换流程。",
        "允许同一套 Sidecar 和 Bridge 框架通过不同 Profile 支持多个柜台。",
        "保留以后将某个 Profile 替换为原生 XTP、TORA 或 QDP API 的边界。",
    ])
    add_heading(doc, "3 2 一期范围", 2)
    add_table(doc, ["范围", "一期内容"], [
        ("行情", "连接、登录、订阅、退订、快照、深度行情、断线重订阅"),
        ("交易", "认证、登录、结算确认、下单、撤单、订单回报、成交回报"),
        ("查询", "账户、持仓、订单、成交、合约及柜台支持的必要查询"),
        ("恢复", "流恢复、活动订单查询、成交补偿、持仓资金对账"),
        ("运维", "健康检查、指标、结构化日志、配置校验、优雅停机和进程监管"),
        ("柜台", "TTS、标准 CTP，并为其余 Profile 建立框架和验收模板"),
    ], widths=[1.15, 5.65])
    add_heading(doc, "3 3 非目标", 2)
    add_bullets(doc, [
        "不在一期内复刻 CTP API 的全部边缘业务接口。未纳入的接口必须显式标记为不支持。",
        "不允许一个 Sidecar 进程同时加载多个同名 OpenCTP 适配库。",
        "不承诺 OpenCTP 转接后的性能等同于柜台原生 API，QDP 等低延迟场景需单独测试。",
        "不使用 TTS 通过作为实盘柜台验收的替代证据。",
        "不在 Connector 内实现策略逻辑、组合风控或账户级资产配置。",
    ])

    page_break(doc)
    add_heading(doc, "4 需求基线", 1)
    add_heading(doc, "4 1 功能需求", 2)
    add_table(doc, ["编号", "需求", "优先级"], [
        ("FR01", "加载指定版本的行情和交易动态库并验证 API 版本", "必须"),
        ("FR02", "管理认证、登录、结算确认、订阅和交易状态机", "必须"),
        ("FR03", "提交、撤销和查询订单并保留原始柜台标识", "必须"),
        ("FR04", "可靠转发订单、成交、账户和持仓事件", "必须"),
        ("FR05", "在断线或重启后完成活动订单、成交和持仓对账", "必须"),
        ("FR06", "按柜台 Profile 处理认证、字段借用和业务差异", "必须"),
        ("FR07", "运行多个账户和柜台实例，每个实例独立配置", "必须"),
        ("FR08", "支持行情与交易进程分离部署", "建议"),
        ("FR09", "支持对单一柜台替换原生 API 而保持 Titan 上层接口不变", "建议"),
    ], widths=[0.72, 5.15, 0.92])
    add_heading(doc, "4 2 非功能需求", 2)
    add_table(doc, ["类别", "要求"], [
        ("正确性", "交易事件不得因队列满而丢弃；重复事件必须幂等处理；未知订单状态必须触发对账"),
        ("隔离", "任一柜台 SO 崩溃不得终止 Titan Core 或其他柜台 Connector"),
        ("可升级", "C ABI 主版本不兼容时拒绝启动；次版本仅允许向后兼容扩展"),
        ("可观测", "连接、请求、回报、队列、延迟、重连、对账和拒单均可度量"),
        ("安全", "凭据不进入日志、命令行和崩溃转储；所有二进制文件校验摘要"),
        ("性能", "Bridge 与 IPC 开销必须独立测量，不能用端到端平均值掩盖尾延迟"),
        ("可测试", "提供 Fake CTP 库、回放模式、故障注入和柜台验收套件"),
    ], widths=[1.15, 5.65])
    add_heading(doc, "4 3 初始工程目标", 2)
    add_para(doc, "下列数值是一期工程目标，不是对第三方柜台的性能承诺。项目组应在目标服务器、目标 SO 和券商测试环境中重新标定。")
    add_table(doc, ["指标", "初始目标", "测量边界"], [
        ("交易事件丢失", "0", "Sidecar 接收至 Titan EventEngine 接纳"),
        ("重复事件", "允许输入重复，最终状态不重复计量", "同一订单和成交的稳定键"),
        ("C ABI 增量开销", "P99 小于 50 微秒", "不含柜台网络与业务处理"),
        ("本机 IPC 增量开销", "P99 小于 100 微秒", "固定负载和目标硬件"),
        ("故障发现", "5 秒内", "Sidecar 退出或心跳中断"),
        ("进程恢复", "60 秒内进入对账或明确失败", "受柜台流控和可用性约束"),
        ("稳定性", "连续 24 小时无未解释内存增长", "行情加交易综合负载"),
    ], widths=[1.5, 2.05, 3.25])

    add_heading(doc, "5 Titan 现状与集成约束", 1)
    add_para(doc, "Titan 当前实盘链路由 titan CLI、TitanCoreRuntime、PluginEngine、EventEngine 与动态 Connector 工厂组成。已有 Binance Futures、OKX 和 Hyperliquid Connector 以 cdylib 插件包提供 MarketConnectorFactory 与 AccountConnectorFactory。插件包带版本和 SHA256 校验，运行时通过 MarketPlugin 和 AccountPlugin 注入事件。")
    add_heading(doc, "5 1 复用能力", 2)
    add_table(doc, ["现有能力", "本方案用法"], [
        ("动态插件入口", "新增 titan connector openctp plugin，沿用 PluginApiV1 和 manifest"),
        ("MarketConnectorFactory", "创建行情代理并将 Sidecar 行情映射到 MarketPlugin"),
        ("AccountConnectorFactory", "创建交易账户代理并将回报映射到 AccountPlugin"),
        ("EventEngine QoS", "交易走 ReliableOrdered，行情走 FastLane 或可丢弃队列"),
        ("包摘要校验", "扩展到 Sidecar、Bridge、OpenCTP SO 和厂商依赖的完整清单"),
        ("RestartRequired", "柜台 profile、动态库或头文件版本变化时重启对应实例"),
    ], widths=[1.7, 5.1])
    add_heading(doc, "5 2 需要新增的能力", 2)
    add_bullets(doc, [
        "插件监管 Sidecar 的启动、心跳、退出、退避重启和优雅停机。",
        "本机 IPC 会话、协议协商和可靠事件确认。",
        "国内柜台 Instrument、OrderRef、FrontID、SessionID、ExchangeID 等标识映射。",
        "CTP 交易日、结算确认、流文件和私有流恢复状态。",
        "柜台 Profile、运行包清单和 ABI 自检。",
    ])

    page_break(doc)
    add_heading(doc, "6 总体架构", 1)
    add_para(doc, "总体架构分为 Titan 插件层、Sidecar 控制层、C ABI 边界、C++ CTP Bridge 和柜台动态库层。Titan Core 不直接加载任何 OpenCTP 或厂商 SO。")
    add_code(doc, r"""
Titan Core Runtime
  MarketPlugin  AccountPlugin  EventEngine
             |
  titan connector openctp plugin
       supervisor and IPC client
             |
       Unix Domain Socket
             |
  titan openctp sidecar process
       state machine and C ABI binding
             |
       titan ctp bridge ABI v1
             |
  C++ CTP API and SPI bridge
       counter profile and mapping
             |
  OpenCTP thost SO and vendor SO
             |
  CTP TTS XTP TORA TAP QDP
""")
    add_para(doc, "图 1 逻辑架构", style="Caption Small")
    add_heading(doc, "6 1 关键边界", 2)
    add_table(doc, ["边界", "允许传递", "禁止传递"], [
        ("Titan 到插件", "Titan Rust trait 与插件 ABI", "CTP 结构体和厂商指针"),
        ("插件到 Sidecar", "版本化 IPC 消息", "裸指针和进程内句柄"),
        ("Rust Sidecar 到 C++ Bridge", "C POD、整数句柄、回调函数表", "C++ 类、异常和 STL 对象"),
        ("Bridge 到 OpenCTP", "对应头文件定义的 CTP API 与 SPI", "不同版本头文件混用"),
    ], widths=[1.45, 2.45, 2.9])
    add_heading(doc, "6 2 数据流", 2)
    add_numbered(doc, [
        "策略或账户插件提交 Titan 订单命令。",
        "OpenCTP 插件分配全局 command id，经可靠 IPC 发送到指定 Sidecar。",
        "Sidecar 写入命令日志并通过 C ABI 调用 C++ Bridge。",
        "Bridge Profile 将 Titan 请求映射为 CThostFtdcInputOrderField 并调用 ReqOrderInsert。",
        "柜台通过 SPI 返回响应、订单和成交事件。",
        "Bridge 在回调线程内复制为固定 POD，立即写入有界通道，不执行阻塞工作。",
        "Sidecar 规范化标识、去重并通过可靠 IPC 发送至插件。",
        "插件发布到 AccountPlugin 和 EventEngine，并向 Sidecar 确认已接纳序号。",
    ])

    add_heading(doc, "7 组件设计", 1)
    add_heading(doc, "7 1 OpenCTP Connector 插件", 2)
    add_para(doc, "插件是 Titan 内部的唯一接入点。它不链接 OpenCTP SO，主要职责是实现 MarketConnectorFactory 和 AccountConnectorFactory、解析配置、启动 Sidecar、协商协议、路由命令、接收事件和报告健康状态。")
    add_bullets(doc, [
        "一个插件包可以创建多个 Connector 实例，但每个交易实例对应独立 Sidecar。",
        "插件不得在 EventEngine 回调线程执行进程启动、文件 IO 或阻塞 IPC。",
        "交易命令进入可靠有序发送队列；行情事件进入独立接收路径。",
        "Sidecar 失联时立即停止接受新交易请求，并向账户插件发布 Degraded 状态。",
    ])
    add_heading(doc, "7 2 Rust Sidecar", 2)
    add_para(doc, "Sidecar 是每个柜台实例的控制平面。它拥有 C ABI handle、会话状态、请求编号、订单映射、事件日志、对账器和健康探针。Sidecar 与 C++ Bridge 位于同一进程，但 Bridge 和第三方 SO 不进入 Titan Core。")
    add_table(doc, ["模块", "职责"], [
        ("SessionController", "驱动连接、认证、登录、结算、查询、就绪和恢复状态"),
        ("CommandRouter", "校验命令状态、限流、生成 request id 并调用 C ABI"),
        ("EventNormalizer", "将 C ABI 事件转换为 Titan 领域事件并保留扩展字段"),
        ("OrderRegistry", "维护 Titan order id 与柜台复合标识映射"),
        ("Reconciler", "查询并收敛订单、成交、持仓和资金状态"),
        ("Journal", "记录关键命令、回报序号和恢复检查点"),
        ("HealthService", "暴露进程、队列、会话、版本和对账健康状态"),
    ], widths=[1.6, 5.2])
    add_heading(doc, "7 3 C++ Bridge", 2)
    add_para(doc, "Bridge 以 C++17 实现，正常继承 CThostFtdcTraderSpi 与 CThostFtdcMdSpi。它持有 API 实例、SPI 实例和 Profile，不将 C++ 异常、字符串、容器或对象指针暴露给 Rust。")
    add_bullets(doc, [
        "所有导出函数均为 extern C，并使用固定宽度整数和显式长度。",
        "C++ 异常在导出函数边界捕获并转换为 TitanCtpStatus。",
        "SPI 回调不得向 C++ 外传播异常，用户回调不得直接重入 CTP API。",
        "API 创建、RegisterSpi、Init、Release 和库卸载在同一所有者线程执行。",
        "Bridge 只复制所需字段；原始 CTP 结构可选择性写入诊断记录，但不跨 ABI。",
    ])
    add_heading(doc, "7 4 Counter Profile", 2)
    add_para(doc, "Profile 是柜台差异的显式策略对象。它提供认证映射、登录补充字段、合约来源、订单字段映射、撤单键、能力声明、错误归类和恢复步骤。Profile 不允许通过散落的 if counter type 修改核心状态机。")

    page_break(doc)
    add_heading(doc, "8 进程与部署模型", 1)
    add_heading(doc, "8 1 进程拓扑", 2)
    add_code(doc, r"""
titan process
  connector openctp plugin
    IPC client ctp account A  ------ sidecar ctp A
    IPC client xtp account B  ------ sidecar xtp B
    IPC client tora account C ------ sidecar tora C

sidecar xtp B process
  rust sidecar
  libtitan ctp bridge xtp.so
  thosttraderapi_se.so
  thostmduserapi_se.so
  xtptraderapi.so and dependencies
""")
    add_para(doc, "图 2 进程隔离与实例关系", style="Caption Small")
    add_heading(doc, "8 2 运行包目录", 2)
    add_code(doc, r"""
/opt/titan/connectors/openctp/xtp/v1/
  manifest.json
  bin/titan-openctp-sidecar
  lib/libtitan_ctp_bridge.so
  vendor/thosttraderapi_se.so
  vendor/thostmduserapi_se.so
  vendor/xtptraderapi.so
  include/ThostFtdcTraderApi.h
  include/ThostFtdcMdApi.h
  include/ThostFtdcUserApiStruct.h
  include/ThostFtdcUserApiDataType.h
  profile/profile.toml
  data/dict.csv
  SHA256SUMS
""")
    add_heading(doc, "8 3 动态库解析", 2)
    add_para(doc, "每个 Sidecar 启动时把自身 vendor 目录放入私有动态库搜索路径。Bridge 通过正常 C++ 链接调用 CTP 类，不在 Rust 中查找 mangled symbol。由于一个进程只加载一个适配器，OpenCTP 各柜台使用相同 thost 库名不会产生进程内冲突。")
    add_para(doc, "运行包不得依赖主机上的全局 LD_LIBRARY_PATH。systemd 或容器入口只为该 Sidecar 设置搜索路径，并在启动前检查所有 DT_NEEDED 依赖。")
    add_heading(doc, "8 4 进程监管", 2)
    add_table(doc, ["场景", "插件行为", "Sidecar 行为"], [
        ("正常启动", "生成会话令牌并等待 Ready", "验证清单、连接、登录并报告阶段"),
        ("心跳超时", "停止新交易并标记 Degraded", "若仍运行则进入自检或退出"),
        ("异常退出", "指数退避重启并触发告警", "重启后强制进入 Reconcile"),
        ("配置更新", "只允许 RestartRequired", "新进程使用新运行包和配置"),
        ("优雅停机", "停止新请求，等待未决命令", "查询活动订单、持久化检查点、Release API"),
    ], widths=[1.2, 2.8, 2.8])

    add_heading(doc, "9 C ABI 设计", 1)
    add_heading(doc, "9 1 ABI 规则", 2)
    add_bullets(doc, [
        "ABI 使用 magic、abi_major、abi_minor、struct_size 和 capability_bits 协商。",
        "主版本不一致时拒绝加载；次版本只允许在结构尾部增加字段和函数指针。",
        "全部结构使用定宽整数、double、字节数组和显式长度，不使用 bool、enum 默认宽度或 size_t。",
        "调用方分配输入内存；事件数据仅在回调期间有效，Sidecar 必须同步复制。",
        "Bridge 不持有 Rust 字符串或切片；所有字符串均带长度并允许非 UTF8 原始字节。",
        "错误通过状态码和线程安全 last_error 返回，禁止异常跨越 ABI。",
    ])
    add_heading(doc, "9 2 ABI 入口", 2)
    add_code(doc, r"""
#define TITAN_CTP_ABI_MAGIC 0x54435450u
#define TITAN_CTP_ABI_MAJOR 1u

typedef struct TitanCtpHandleImpl* TitanCtpHandle;

typedef struct TitanCtpApiV1 {
  uint32_t magic;
  uint16_t abi_major;
  uint16_t abi_minor;
  uint32_t struct_size;
  uint64_t capability_bits;

  TitanCtpStatus (*create)(const TitanCtpConfigV1*,
                           const TitanCtpCallbacksV1*,
                           TitanCtpHandle*);
  TitanCtpStatus (*start)(TitanCtpHandle);
  TitanCtpStatus (*authenticate)(TitanCtpHandle,
                                 const TitanCtpAuthV1*);
  TitanCtpStatus (*login)(TitanCtpHandle,
                          const TitanCtpLoginV1*);
  TitanCtpStatus (*submit_order)(TitanCtpHandle,
                                 const TitanOrderRequestV1*);
  TitanCtpStatus (*cancel_order)(TitanCtpHandle,
                                 const TitanCancelRequestV1*);
  TitanCtpStatus (*query)(TitanCtpHandle,
                          const TitanQueryRequestV1*);
  TitanCtpStatus (*stop)(TitanCtpHandle, uint32_t timeout_ms);
  void (*destroy)(TitanCtpHandle);
  uint32_t (*last_error)(TitanCtpHandle, char*, uint32_t);
} TitanCtpApiV1;

extern "C" const TitanCtpApiV1* titan_ctp_bridge_entry_v1(void);
""")
    add_heading(doc, "9 3 结构版本策略", 2)
    add_table(doc, ["字段", "用途"], [
        ("struct_size", "调用方填写自身可用字节数，Bridge 只访问已声明范围"),
        ("schema_version", "领域载荷版本，与 ABI 函数表版本分离"),
        ("flags", "定义可选行为，未知 bit 必须拒绝或忽略并记录"),
        ("extensions_ptr", "指向 TLV 扩展数据，仅在调用期间有效"),
        ("reserved", "保持零值，为兼容扩展预留"),
    ], widths=[1.45, 5.35])
    add_heading(doc, "9 4 回调函数表", 2)
    add_code(doc, r"""
typedef struct TitanCtpCallbacksV1 {
  uint32_t struct_size;
  void* context;
  void (*on_state)(void*, const TitanSessionEventV1*);
  void (*on_market_data)(void*, const TitanMarketEventV1*);
  void (*on_order)(void*, const TitanOrderEventV1*);
  void (*on_trade)(void*, const TitanTradeEventV1*);
  void (*on_query_row)(void*, const TitanQueryRowV1*);
  void (*on_error)(void*, const TitanErrorEventV1*);
} TitanCtpCallbacksV1;
""")
    add_para(doc, "C++ SPI 回调先复制数据，再调用上述函数。Rust 回调只能把事件写入 Sidecar 内部队列，不得等待网络、锁住 API 所有者线程或抛出 panic。Rust 导出回调统一使用 catch_unwind，发生 panic 时记录错误并要求 Sidecar 退出。")

    page_break(doc)
    add_heading(doc, "10 IPC 与事件传输", 1)
    add_heading(doc, "10 1 通道划分", 2)
    add_table(doc, ["通道", "内容", "QoS", "背压"], [
        ("控制通道", "握手、配置、健康、停止、查询", "可靠有序", "阻塞发送线程，不阻塞 CTP 回调"),
        ("交易命令", "下单、撤单、查询", "可靠有序", "达到上限时拒绝新命令"),
        ("交易事件", "订单、成交、错误、账户、持仓", "可靠有序", "落盘后重试，禁止丢弃"),
        ("行情事件", "Tick、盘口、订阅状态", "尽力有序", "按合约合并最新值并计数"),
    ], widths=[1.15, 2.25, 1.2, 2.2])
    add_heading(doc, "10 2 一期协议", 2)
    add_para(doc, "一期使用 Unix Domain Socket 的长度前缀二进制帧。协议载荷采用稳定的显式 schema，建议使用 Protobuf 或 FlatBuffers；若沿用 Titan 现有 Rust 编码，必须提供 C++ 可生成的同源 schema，不能手写两套结构。交易事件启用 sequence 和 ack，行情不要求逐条 ack。")
    add_code(doc, r"""
FrameHeaderV1
  magic          u32
  protocol_major u16
  protocol_minor u16
  message_type   u32
  flags          u32
  payload_len    u32
  session_id     u64
  sequence       u64
  correlation_id u64
  crc32c         u32
""")
    add_heading(doc, "10 3 序号和确认", 2)
    add_bullets(doc, [
        "Sidecar 为交易事件分配单调递增 sequence，并在 Journal 保存已发送范围。",
        "插件只在事件成功进入 Titan 可靠队列后发送 ack。",
        "重连后 Sidecar 从 last_acked_sequence 重放尚未确认的事件。",
        "Titan 依据事件稳定键去重，因此重放不会重复增加成交和持仓。",
        "行情通道独立编号，用于监控缺口，但默认不重放历史 Tick。",
    ])
    add_heading(doc, "10 4 后续低延迟路径", 2)
    add_para(doc, "当 QDP 或高频行情证明 UDS 不能满足尾延迟目标时，可把数据面替换为共享内存 SPSC ring，控制面继续使用 UDS。该优化不改变 C ABI、IPC 消息 schema 或 Titan Connector 接口。ring 满时，交易通道触发熔断；行情通道按合约覆盖最旧快照并增加 drop 和 coalesce 指标。")

    add_heading(doc, "11 领域模型与标识", 1)
    add_heading(doc, "11 1 标识体系", 2)
    add_table(doc, ["标识", "来源", "用途"], [
        ("command_id", "Titan 插件", "命令幂等和 IPC 关联"),
        ("order_id", "Titan", "上层稳定订单主键"),
        ("request_id", "Sidecar", "CTP 请求响应关联，交易日内递增"),
        ("order_ref", "Sidecar 或柜台", "CTP 下单引用，Profile 管理格式和序列"),
        ("front_id session_id", "柜台登录响应", "撤单与订单定位"),
        ("order_sys_id", "交易所或柜台", "最终柜台订单标识"),
        ("trade_id", "交易所或柜台", "成交去重"),
        ("instrument_key", "ExchangeID 加 InstrumentID", "跨市场唯一合约键"),
    ], widths=[1.55, 1.7, 3.55])
    add_heading(doc, "11 2 订单状态映射", 2)
    add_para(doc, "状态映射必须保留事实和推导的区别。CTP 的 OrderStatus、OrderSubmitStatus、VolumeTraded、VolumeTotal 与错误回调共同决定 Titan 状态。Connector 不因收到成功的 ReqOrderInsert 返回值就标记订单已接受，该返回值只说明请求进入客户端 API。")
    add_table(doc, ["CTP 事实", "Titan 状态", "处理"], [
        ("ReqOrderInsert 返回非零", "RejectedLocal", "记录本地拒绝，不等待回报"),
        ("OnRspOrderInsert 错误", "Rejected", "保存 ErrorID 与原始信息"),
        ("OnRtnOrder 已报", "Accepted 或 Working", "记录柜台复合键"),
        ("部分成交", "PartiallyFilled", "分别发布 OrderEvent 和每笔 TradeEvent"),
        ("全部成交", "Filled", "以累计数量核对成交明细"),
        ("已撤单", "Canceled", "保留已成交数量"),
        ("状态未知或断线", "PendingReconcile", "停止推测并执行查询"),
    ], widths=[2.0, 1.45, 3.35])
    add_heading(doc, "11 3 金额与数量", 2)
    add_para(doc, "CTP 使用 double 表示价格和金额，Titan 当前 OrderEvent 与 FillEvent 同样使用 f64。Connector 可以无损传递底层 bit pattern，但下单前仍要按合约 tick size、volume multiple 和最小数量验证。资金报表和对账结果建议同时保存原始 double 与标准化十进制定点值，避免跨系统汇总误差。")
    add_heading(doc, "11 4 扩展字段", 2)
    add_para(doc, "扩展区采用版本化 TLV 或 schema oneof，保存 shareholder account、market、hedge flag、covered flag、ClientID、business unit、front/session、原始错误码等字段。未知扩展必须透传或记录，不能静默删除。")

    page_break(doc)
    add_heading(doc, "12 生命周期与恢复", 1)
    add_heading(doc, "12 1 会话状态机", 2)
    add_code(doc, r"""
Created
  -> ValidatingPackage
  -> Connecting
  -> Authenticating
  -> LoggingIn
  -> ConfirmingSettlement
  -> LoadingReferenceData
  -> Reconciling
  -> Ready
  -> Degraded
  -> Reconnecting
  -> Reconciling
  -> Ready
  -> Stopping
  -> Stopped
""")
    add_para(doc, "图 3 Sidecar 会话状态机", style="Caption Small")
    add_heading(doc, "12 2 启动流程", 2)
    add_numbered(doc, [
        "验证 manifest、SHA256、文件权限、动态库依赖、Profile 和凭据引用。",
        "加载 Bridge，校验 C ABI、capabilities、编译头文件版本和运行时 GetApiVersion。",
        "创建行情和交易 API，注册 SPI，设置订阅流模式并调用 Init。",
        "完成认证和登录，保存 TradingDay、FrontID、SessionID 和最大 OrderRef。",
        "按 Profile 执行结算确认、合约加载或外部字典检查。",
        "查询活动订单、成交、持仓和资金，与 Journal 和 Titan 状态对账。",
        "对账完成后进入 Ready，并允许新交易命令。",
    ])
    add_heading(doc, "12 3 断线恢复", 2)
    add_table(doc, ["阶段", "动作"], [
        ("发现断线", "发布 Degraded，停止接受新下单，可按策略允许撤单进入待发队列"),
        ("重新连接", "遵守柜台退避和限流，避免多个实例同步重连"),
        ("重新登录", "更新 FrontID、SessionID、最大 OrderRef 和交易日"),
        ("恢复私有流", "根据 QUICK、RESUME、RESTART 和 Profile 能力选择"),
        ("查询补偿", "查询订单、成交、持仓、资金，补发缺失事件"),
        ("状态收敛", "所有未知订单得到柜台事实后恢复 Ready"),
    ], widths=[1.45, 5.35])
    add_heading(doc, "12 4 交易日切换", 2)
    add_bullets(doc, [
        "TradingDay 改变时冻结旧交易日 OrderRef 分配器并创建新命名空间。",
        "结算确认按柜台要求重新执行，完成前不得开放交易。",
        "流文件按账户、柜台、环境、交易日和 API 类型分目录保存。",
        "跨夜活动订单按交易所事实处理，不能统一假设已撤销。",
        "合约、涨跌停、手续费和保证金等参考数据按交易日刷新。",
    ])
    add_heading(doc, "12 5 优雅停机", 2)
    add_para(doc, "Titan 当前交易连接具备退出前取消未完成委托的安全思路。国内柜台未必提供统一 cancel all，Sidecar 应先停止新请求，查询活动订单，再按配置逐笔撤单并等待回报。若柜台不可用，必须报告未确认订单清单，不能把进程退出等同于撤单成功。")

    add_heading(doc, "13 柜台 Profile", 1)
    add_heading(doc, "13 1 能力声明", 2)
    add_code(doc, r"""
profile.id = "xtp"
profile.asset_classes = ["stock", "fund", "bond", "etf_option"]
profile.supports.instrument_query = false
profile.supports.settlement_confirm = false
profile.supports.private_stream_resume = true
profile.requires.instrument_dictionary = true
profile.cancel_key = "profile_defined"
profile.auth_mapping = "xtp_openctp_v1"
""")
    add_heading(doc, "13 2 Profile 差异矩阵", 2)
    add_table(doc, ["Profile", "主要资产", "必须处理的差异", "风险"], [
        ("ctp", "期货和期货期权", "认证、结算、开平和平今、交易日", "中"),
        ("tts", "模拟股票期货期权", "模拟撮合、版本与环境能力", "低"),
        ("ctp sopt", "ETF 股票期权", "SOPT namespace、行权、备兑和持仓", "高"),
        ("xtp", "股票债券基金 ETF期权", "长授权码、ClientID、dict.csv、证券账户", "高"),
        ("tora", "股票债券基金", "TerminalInfo、FENS、OrderRef 与 ActionRef", "高"),
        ("tora opt", "股票期权", "期权交易 API、行情共享、备兑与行权", "高"),
        ("tap", "期货及相关品种", "合约编码、订单类型、开平与查询限流", "中"),
        ("qdp", "期货", "版本、回报时序和尾延迟", "中高"),
    ], widths=[0.9, 1.55, 3.55, 0.8], font_size=8.2)
    add_heading(doc, "13 3 XTP 特别设计", 2)
    add_bullets(doc, [
        "授权码超过标准 CTP AuthCode 长度时，禁止在 Rust 或 C++ 中依赖未定义的数组越界。应要求 OpenCTP 提供明确配置入口，或在受控 Bridge 内按经验证的适配器协议构造专用认证缓冲区。",
        "dict.csv 作为运行包受管资产，包含版本、生成时间、市场和 SHA256；交易日前校验新鲜度。",
        "ClientID 不得硬编码为进程级常量。若适配器固定使用某值，部署系统必须检测冲突并禁止重复实例。",
        "股票 T 加一、股东账户、市场和证券业务类型必须进入 Profile 与扩展字段。",
    ])
    add_heading(doc, "13 4 TORA 特别设计", 2)
    add_bullets(doc, [
        "集中管理 OrderRef 与 OrderActionRef 共享序列，任何调用方不得自行填写。",
        "Profile 构造 TerminalInfo，并对 IP、MAC、硬盘序列等来源和脱敏策略进行配置。",
        "FENS 地址中的 EnvID 和 NodeID 由结构化配置生成，不允许策略提供原始拼接字符串。",
        "股票、两融和股票期权能力分开声明，不能用同一 profile 暗示全量业务支持。",
    ])
    add_heading(doc, "13 5 TAP 与 QDP 特别设计", 2)
    add_para(doc, "TAP 与 QDP 的期货语义更接近 CTP，但仍需针对交易所编码、组合、条件单、FAK、FOK、平今、限流和回报顺序逐项验证。QDP Profile 另外启用低开销日志和性能采样，并保留切换原生 QDP Connector 的上层接口。")

    page_break(doc)
    add_heading(doc, "14 性能设计", 1)
    add_heading(doc, "14 1 热路径", 2)
    add_code(doc, r"""
CTP SPI thread
  copy selected fields to POD
  timestamp with monotonic clock
  enqueue to bounded lock free queue
  return immediately

Sidecar event thread
  normalize and deduplicate
  append reliable trading event
  encode IPC frame
  send to Titan plugin
""")
    add_heading(doc, "14 2 性能控制", 2)
    add_bullets(doc, [
        "行情和交易使用不同队列、线程和指标，行情高峰不得阻塞成交回报。",
        "SPI 回调不做磁盘 IO、DNS、日志格式化、锁等待或 GB18030 大对象转换。",
        "高频字段使用固定数组和预分配 slab，避免每 Tick heap allocation。",
        "原始接收时间、Bridge 入队时间、Sidecar 发送时间和 Titan 接纳时间全部打点。",
        "QDP 验收同时比较 OpenCTP 转接路径与厂商原生 API，报告 P50、P99、P99 9 和最大值。",
    ])
    add_heading(doc, "14 3 背压和内存", 2)
    add_table(doc, ["队列", "满载策略"], [
        ("交易命令", "拒绝新命令并返回明确的 ConnectorBusy，不覆盖旧命令"),
        ("交易事件", "触发交易熔断，保留事件到 Journal，禁止丢弃"),
        ("查询结果", "按 request id 流式发送，必要时暂停新的非关键查询"),
        ("行情", "同合约最新值覆盖，保留 drop 和 coalesce 计数"),
        ("日志", "异步批量写入，错误日志保留独立容量"),
    ], widths=[1.45, 5.35])

    add_heading(doc, "15 安全设计", 1)
    add_heading(doc, "15 1 凭据", 2)
    add_bullets(doc, [
        "配置文件只保存 secret reference，凭据通过受限文件描述符、密钥服务或受控环境注入。",
        "用户名、密码、AuthCode、AppID 和终端信息不得出现在进程参数、普通日志和 core dump。",
        "Bridge 使用后清零临时密码缓冲区；错误消息不得回显完整认证字段。",
        "生产与仿真凭据、库目录、流文件和网络策略完全分离。",
    ])
    add_heading(doc, "15 2 二进制供应链", 2)
    add_table(doc, ["控制", "要求"], [
        ("来源", "记录 OpenCTP 与厂商下载来源、时间、版本和审批人"),
        ("完整性", "Sidecar 启动前验证 SHA256SUMS，插件包验证扩展清单"),
        ("可追溯", "运行日志打印非敏感 build id、API version 和 manifest digest"),
        ("权限", "vendor 目录只读，Sidecar 使用非特权用户运行"),
        ("网络", "只允许访问配置的行情交易前置和本机 IPC"),
        ("升级", "新版本进入隔离验证环境，不能原地覆盖生产目录"),
    ], widths=[1.25, 5.55])
    add_heading(doc, "15 3 C 与 C++ 安全边界", 2)
    add_bullets(doc, [
        "使用 AddressSanitizer、UndefinedBehaviorSanitizer 和 ThreadSanitizer 测试 Bridge 与 Fake CTP。",
        "对所有长度、数组和字符串终止符进行校验，禁止依赖结构相邻布局实现越界协议。",
        "每个导出函数捕获所有 C++ 异常；析构函数不得抛出。",
        "所有 callback context 在 destroy 前撤销注册并等待回调线程退出。",
    ])

    add_heading(doc, "16 可观测性与运维", 1)
    add_heading(doc, "16 1 指标", 2)
    add_table(doc, ["域", "指标示例"], [
        ("会话", "connection_state、login_total、reconnect_total、trading_day"),
        ("命令", "submit_total、cancel_total、local_reject_total、request_rate_limited_total"),
        ("回报", "order_event_total、trade_event_total、duplicate_total、sequence_gap_total"),
        ("队列", "depth、capacity、high_watermark、drop_total、coalesce_total"),
        ("延迟", "abi_latency、ipc_latency、order_round_trip、callback_to_engine"),
        ("恢复", "reconcile_total、reconcile_duration、unknown_order_count"),
        ("进程", "restart_total、rss_bytes、cpu、fd_count、heartbeat_age"),
    ], widths=[1.1, 5.7])
    add_heading(doc, "16 2 日志", 2)
    add_para(doc, "所有日志使用结构化字段，并包含 connector_instance、profile、environment、account_alias、trading_day、session_id、command_id、request_id 和 order_id。密码、AuthCode、完整账户号和原始终端指纹必须脱敏。")
    add_heading(doc, "16 3 告警", 2)
    add_bullets(doc, [
        "Sidecar 退出、重启循环或心跳超过阈值。",
        "交易事件队列达到高水位或出现 sequence gap。",
        "登录失败、认证失败、结算未确认或交易日异常。",
        "对账后出现未知活动订单、成交缺口或持仓资金不一致。",
        "动态库摘要、API 版本或 Profile 能力不匹配。",
        "拒单率、撤单失败率或端到端延迟显著偏离基线。",
    ])
    add_heading(doc, "16 4 运维操作", 2)
    add_table(doc, ["操作", "控制"], [
        ("启动", "先验证包和网络，再连接，不允许跳过对账进入 Ready"),
        ("停机", "停止新请求，处理未决命令，查询活动订单并持久化检查点"),
        ("升级", "新目录部署、仿真验收、灰度账户、可回退到旧目录"),
        ("故障切换", "新实例必须使用不同 ClientID 与流目录，先对账后接管"),
        ("人工处置", "所有强制 Ready、忽略差异等操作需要审计记录，生产默认禁用"),
    ], widths=[1.2, 5.6])

    page_break(doc)
    add_heading(doc, "17 测试与验收", 1)
    add_heading(doc, "17 1 测试分层", 2)
    add_table(doc, ["层级", "重点"], [
        ("单元测试", "字段映射、状态机、标识、错误分类、Profile 规则"),
        ("ABI 合同测试", "版本协商、结构大小、空指针、长度、异常和释放"),
        ("Fake CTP", "同步回调、异步回调、乱序、重复、断线、崩溃和限流"),
        ("TTS 集成", "登录、行情、下单、撤单、部分成交、查询和重连"),
        ("柜台仿真", "XTP、TORA、TAP、QDP 各自认证和语义"),
        ("压力测试", "行情突发、订单突发、队列满、长时间运行和内存"),
        ("故障演练", "kill sidecar、网络分区、SO 崩溃、磁盘满和交易日切换"),
        ("实盘验收", "小额白名单、撤单保护、最终柜台查询对账"),
    ], widths=[1.25, 5.55])
    add_heading(doc, "17 2 Fake CTP 动态库", 2)
    add_para(doc, "测试仓库应提供与目标头文件匹配的 Fake CTP SO。它导出相同 Create 与 GetApiVersion 符号，并允许脚本控制 SPI 回调。Fake 库用于稳定复现真实柜台难以制造的竞态和故障。")
    add_bullets(doc, [
        "ReqOrderInsert 返回前同步触发回调。",
        "订单回报与成交回报交换顺序。",
        "重复成交、空指针响应和非 UTF8 错误消息。",
        "Release 期间仍有回调在飞行。",
        "断线后切换 FrontID、SessionID 和 TradingDay。",
        "查询多行结果中断、缺少 isLast 或返回流控错误。",
    ])
    add_heading(doc, "17 3 柜台验收矩阵", 2)
    add_table(doc, ["能力", "CTP", "TTS", "XTP", "TORA", "TAP", "QDP"], [
        ("认证登录", "验", "验", "专项", "专项", "专项", "专项"),
        ("行情订阅", "验", "验", "验", "验", "验", "验"),
        ("普通下单撤单", "验", "验", "验", "验", "验", "验"),
        ("部分成交", "验", "验", "验", "验", "验", "验"),
        ("订单成交查询", "验", "验", "专项", "验", "验", "验"),
        ("持仓资金", "验", "验", "验", "验", "验", "验"),
        ("断线恢复", "验", "验", "验", "验", "验", "验"),
        ("柜台特有能力", "期权", "模拟", "字典账户", "终端信息", "订单类型", "延迟"),
    ], widths=[1.55, 0.7, 0.7, 0.85, 0.85, 0.85, 0.85], font_size=7.9)
    add_heading(doc, "17 4 生产准入条件", 2)
    add_bullets(doc, [
        "目标动态库、头文件和依赖已形成签名运行包，并通过 ABI 自检。",
        "完整交易回报零丢失，重复注入后账户和成交结果保持幂等。",
        "断线、Sidecar 崩溃和 Titan 插件重连后均能通过查询收敛。",
        "24 小时压力运行无未解释内存增长、线程增长或文件描述符泄漏。",
        "关键延迟达到该柜台批准的基线，QDP 完成原生 API 对照。",
        "实盘小额验收完成，最终柜台查询与 Titan 状态一致。",
        "运维手册、监控、告警、回退包和应急联系人齐备。",
    ])

    add_heading(doc, "18 实施计划", 1)
    add_heading(doc, "18 1 阶段与交付物", 2)
    add_table(doc, ["阶段", "工作", "交付物", "估算"], [
        ("阶段零", "版本和 ABI Spike", "TTS 加载、版本校验、最小回调", "1 至 2 周"),
        ("阶段一", "C ABI 与 Fake CTP", "bridge v1、代码生成、合同测试", "2 至 3 周"),
        ("阶段二", "Sidecar 与 IPC", "监管、协议、Journal、健康检查", "3 至 4 周"),
        ("阶段三", "Titan 插件集成", "Market 与 Account 工厂、事件映射", "2 至 3 周"),
        ("阶段四", "TTS 与官方 CTP", "完整交易链路和恢复验收", "3 至 4 周"),
        ("阶段五", "柜台 Profile", "每个柜台能力和验收包", "每个 2 至 4 周"),
        ("阶段六", "生产加固", "压力、安全、运维、灰度和回退", "3 至 4 周"),
    ], widths=[0.8, 1.65, 3.2, 1.15], font_size=8.3)
    add_para(doc, "估算按具备 Rust、C++ 和柜台联调经验的小团队计算。多个 Profile 可在基础框架稳定后并行推进，实际日历时间取决于测试账号、柜台环境、券商支持和接口包获取。")
    add_heading(doc, "18 2 建议仓库结构", 2)
    add_code(doc, r"""
crates/titan-connector-openctp-plugin/
crates/titan-openctp-sidecar/
crates/titan-openctp-protocol/
crates/titan-openctp-ffi/
native/titan-ctp-bridge/
  include/titan_ctp_bridge_v1.h
  src/api_session.cpp
  src/spi_bridge.cpp
  src/profile.cpp
  profiles/ctp.cpp
  profiles/tts.cpp
  profiles/xtp.cpp
  profiles/tora.cpp
  profiles/tap.cpp
  profiles/qdp.cpp
tests/fake-ctp/
tests/openctp-acceptance/
packaging/openctp/
""")
    add_heading(doc, "18 3 代码生成", 2)
    add_para(doc, "不建议手写数百个 CTP 结构布局。构建工具从指定头文件生成 C POD 转换代码、字段清单和 static_assert。生成物必须带头文件摘要，并在代码评审中与版本变更一并提交。ctp2rs 的 AST 解析和版本分支可作为实现参考，但生成目标是 C++ 内部映射和小型 C ABI，不是 Rust C++ 虚表。")
    add_heading(doc, "18 4 团队职责", 2)
    add_table(doc, ["角色", "职责"], [
        ("架构负责人", "冻结 C ABI、IPC、状态机和版本策略"),
        ("C++ 工程", "Bridge、SPI 生命周期、Profile 和 sanitizer"),
        ("Rust 工程", "Sidecar、IPC、Titan 插件、对账和指标"),
        ("柜台工程", "账号环境、业务语义、错误码和验收场景"),
        ("测试工程", "Fake CTP、压力、故障注入和回归矩阵"),
        ("运维安全", "运行包、凭据、部署、监控、灰度和应急"),
    ], widths=[1.35, 5.45])

    page_break(doc)
    add_heading(doc, "19 风险与决策", 1)
    add_heading(doc, "19 1 风险登记", 2)
    add_table(doc, ["风险", "影响", "控制"], [
        ("头文件与 SO 不匹配", "崩溃或静默数据错位", "版本清单、摘要、GetApiVersion、static_assert 和启动拒绝"),
        ("闭源适配器缺陷", "Sidecar 崩溃或错误回报", "进程隔离、监管重启、Journal 和对账"),
        ("交易事件背压", "丢失订单或成交事实", "可靠队列、落盘、熔断，禁止覆盖"),
        ("股票语义被 CTP 抽象丢失", "错误下单或持仓解释", "Profile、扩展字段和柜台验收"),
        ("XTP 认证特殊约定", "无法登录或内存越界", "禁止通用越界方案，与适配器方确认受支持入口"),
        ("多实例 ClientID 冲突", "登录互踢或请求混淆", "部署注册表、启动检测和账户级锁"),
        ("QDP 转接延迟", "低延迟目标落空", "原生对照基准和保留原生 Connector 路径"),
        ("零自动化柜台环境", "升级回归不足", "Fake CTP 加 TTS 加柜台日常探针"),
    ], widths=[1.8, 2.0, 3.0], font_size=8.2)
    add_heading(doc, "19 2 待决事项", 2)
    add_table(doc, ["事项", "建议", "冻结时点"], [
        ("IPC schema", "在 Protobuf 与 FlatBuffers 中选择并做延迟原型", "阶段零结束"),
        ("Journal 技术", "二进制 append only WAL，按交易日分段", "阶段一结束"),
        ("行情进程是否分离", "先同 Sidecar 双队列，压力不足再拆分", "TTS 压测后"),
        ("XTP 授权码接口", "取得 OpenCTP 书面或可测试约定后实现", "XTP 开发前"),
        ("C ABI 字段范围", "以 Titan 一期事件模型和柜台矩阵冻结", "阶段一评审"),
        ("生产自动撤单", "按柜台能力和业务策略配置，默认逐笔核实", "实盘验收前"),
    ], widths=[1.55, 3.8, 1.45])
    add_heading(doc, "19 3 架构决策记录", 2)
    add_table(doc, ["编号", "决策", "状态"], [
        ("ADR001", "第三方柜台库运行在独立 Sidecar 进程", "建议采纳"),
        ("ADR002", "C++ Bridge 暴露版本化 C ABI，不暴露 C++ 类型", "建议采纳"),
        ("ADR003", "每个 Sidecar 只加载一个柜台适配器实例族", "建议采纳"),
        ("ADR004", "交易与行情采用独立 QoS 和队列", "建议采纳"),
        ("ADR005", "ctp2rs 用作版本与 API 参考，不作为生产 ABI 层", "建议采纳"),
        ("ADR006", "柜台差异通过 Profile 和扩展字段表达", "建议采纳"),
    ], widths=[0.85, 4.85, 1.1])

    add_heading(doc, "20 附录", 1)
    add_heading(doc, "20 1 配置示例", 2)
    add_code(doc, r"""
connector_type = "openctp"
instance_id = "xtp-account-a"
profile = "xtp"
environment = "production"

[process]
sidecar = "/opt/titan/connectors/openctp/xtp/v1/bin/titan-openctp-sidecar"
runtime_dir = "/var/lib/titan/openctp/xtp-account-a"
restart_policy = "on-failure"

[library]
bridge = "lib/libtitan_ctp_bridge.so"
vendor_dir = "vendor"
manifest = "manifest.json"

[front]
trader = "secret-or-config-reference"
market = "secret-or-config-reference"

[credentials]
reference = "vault://trading/xtp/account-a"

[limits]
command_queue = 4096
trading_event_queue = 65536
market_event_queue = 262144
""")
    add_heading(doc, "20 2 Manifest 示例", 2)
    add_code(doc, r"""
{
  "package_schema": 1,
  "profile": "xtp",
  "bridge_abi": {"major": 1, "minor": 0},
  "ctp_header_version": "6.7.x-profile-approved",
  "expected_runtime_api_version": "profile-approved-pattern",
  "files": [
    {"path": "lib/libtitan_ctp_bridge.so", "sha256": "..."},
    {"path": "vendor/thosttraderapi_se.so", "sha256": "..."},
    {"path": "vendor/thostmduserapi_se.so", "sha256": "..."},
    {"path": "data/dict.csv", "sha256": "..."}
  ]
}
""")
    add_heading(doc, "20 3 错误分类", 2)
    add_table(doc, ["类别", "示例", "处理"], [
        ("Configuration", "缺少文件、版本不匹配、ClientID 冲突", "拒绝启动"),
        ("Authentication", "认证码、AppID、密码错误", "停止重试并告警"),
        ("Transport", "前置断线、超时、网络错误", "退避重连并对账"),
        ("FlowControl", "查询或请求过快", "按柜台规则延迟重试"),
        ("BusinessReject", "价格、数量、权限、资金拒绝", "映射到订单拒绝"),
        ("Protocol", "未知回调、结构不完整、序号缺口", "降级并对账"),
        ("Internal", "Bridge 异常、队列满、Journal 失败", "交易熔断或退出"),
    ], widths=[1.25, 2.65, 2.9])
    add_heading(doc, "20 4 参考资料", 2)
    refs = [
        ("OpenCTP 项目和柜台兼容接口", "https://github.com/openctp/openctp"),
        ("ctp2rs 动态加载与 OpenCTP 示例", "https://github.com/pseudocodes/ctp2rs"),
        ("Titan Connector 现状", str(ROOT / "connector/README.md")),
        ("Titan Runtime ABI", str(ROOT / "crates/titan-runtime-abi/src/lib.rs")),
        ("Titan Connector Loader", str(ROOT / "crates/titan-connector-loader/src/lib.rs")),
    ]
    for text, url in refs:
        p = doc.add_paragraph(style="List Bullet")
        if url.startswith("http"):
            add_hyperlink(p, text, url)
        else:
            r = p.add_run(f"{text}  {url}")
            set_east_asia(r)
    add_heading(doc, "20 5 首次评审清单", 2)
    add_bullets(doc, [
        "确认进程分层是否符合 Titan 部署和运维方式。",
        "冻结一期必须支持的 CTP 请求、SPI 回调和柜台能力。",
        "确定 IPC schema 和 Journal 技术选型。",
        "取得 TTS、官方 CTP 和首个生产柜台的版本化接口包与测试账号。",
        "确认 XTP 授权码、ClientID 和合约字典的受支持接入方式，并批准阶段零 Spike 与生产准入指标。",
    ])

    configure_header_footer(doc)
    doc.core_properties.title = "Titan OpenCTP Connector 技术方案"
    doc.core_properties.subject = "自建薄 C++ Connector 与版本化 C ABI"
    doc.core_properties.author = "Titan 工程团队"
    doc.core_properties.keywords = "Titan OpenCTP CTP Connector C ABI Sidecar"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
